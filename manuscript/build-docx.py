import re, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml.ns import qn

src_dir, out = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
pf = st.paragraph_format
pf.left_indent = Inches(0); pf.right_indent = Inches(0)
pf.space_before = Pt(0); pf.space_after = Pt(0)
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = 1.15
pf.first_line_indent = Inches(0.5)
for s in doc.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)

tok = re.compile(r'(\*\*.+?\*\*|\*.+?\*)')
def runs(p, text):
    for part in tok.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 1:
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)

def centered(text, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text); r.bold = bold; r.italic = italic
    return p

first = True
for i in range(1, 10):
    text = open(f'{src_dir}/ch0{i}.md', encoding='utf-8').read()
    blocks = [b.strip() for b in text.split('\n\n') if b.strip()]
    for b in blocks:
        b = ' '.join(b.split('\n'))
        if b.startswith('# '):
            p = centered(b[2:], bold=True)
            if not first:
                p.paragraph_format.page_break_before = True
            first = False
        elif re.fullmatch(r'\*[^*]+\*', b) and '—' in b and blocks.index(b) < 3:
            centered(b[1:-1], italic=True)
            doc.add_paragraph().paragraph_format.first_line_indent = Inches(0)
        elif b == '#':
            centered('#')
        else:
            runs(doc.add_paragraph(), b)
doc.core_properties.title = 'Twenty-Seven Summers in Tuscany — Chapters 1–9'
doc.save(out)
